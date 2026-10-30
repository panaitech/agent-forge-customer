# 读取人维护的价格表。核算模板不存单价，报价时按条件来取。
# 不读取「供应商」表。不维护「是否当前」「生效日期」。
# 模组一行一规格；单面/双面含税单价同属一行两列，按核算模板「灌胶」选用。

from openpyxl import load_workbook


def _rows(ws, start=5):
    # 第 4 行是表头，数据从第 5 行起。空行跳过。
    head = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    out = []
    for raw in ws.iter_rows(min_row=start, values_only=True):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        item = {}
        for i, key in enumerate(head):
            if key:
                item[str(key)] = raw[i] if i < len(raw) else None
        out.append(item)
    return out


def _num(v):
    if v in (None, ""):
        return None
    return float(v)


def _near(a, b):
    if a is None or b is None:
        return False
    return abs(float(a) - float(b)) < 0.001


def unit_of(row, glue):
    # 按灌胶面别取同行的单面或双面含税单价。
    if glue == "双面":
        return _num(row.get("含税单价_双面"))
    return _num(row.get("含税单价_单面"))


def load(path):
    # data_only=False：含税单价在程序里按同一公式重算，不依赖表格是否已经打开过。
    wb = load_workbook(path, data_only=False)
    lamp = _rows(wb["灯珠单价"])
    kit = _rows(wb["套件单价"])
    box = _rows(wb["箱体单价"])
    mod = _rows(wb["模组价格"])
    tpl = _rows(wb["核算模板"])
    for row in kit:
        shell = _num(row.get("底壳单价")) or 0
        mask_n = _num(row.get("面罩数量")) or 0
        mask = _num(row.get("面罩单价")) or 0
        ring = _num(row.get("胶圈单价")) or 0
        row["套件合计"] = shell + mask_n * mask + ring
    for row in mod:
        batch = row.get("灯珠批次")
        lamps = [x for x in lamp if x.get("灯珠批次") == batch]
        one = lamps[0] if len(lamps) == 1 else None
        rq = _num(row.get("R数量")) or 0
        gq = _num(row.get("G数量")) or 0
        bq = _num(row.get("B数量")) or 0
        rp = _num(one.get("R单价")) if one else 0
        gp = _num(one.get("G单价")) if one else 0
        bp = _num(one.get("B单价")) if one else 0
        bead = rq * (rp or 0) + gq * (gp or 0) + bq * (bp or 0)
        kits = [
            x for x in kit
            if _near(x.get("点间距_mm"), row.get("点间距_mm")) and x.get("工艺") == row.get("工艺")
        ]
        kit_price = kits[0]["套件合计"] if len(kits) == 1 else 0
        glue = _num(row.get("灌胶加价")) or 0
        fee = _num(row.get("加工费")) or 0
        single = bead + fee + kit_price
        row["灯珠合计"] = bead
        row["套件单价"] = kit_price
        row["含税单价_单面"] = single
        row["含税单价_双面"] = single + glue
        row["灯珠单价行数"] = len(lamps)
        row["套件行数"] = len(kits)
    return {"mod": mod, "box": box, "tpl": tpl}


def _pitch(v):
    n = float(v)
    return str(int(n)) if n == int(n) else f"{n:g}"


def mod_spec(row, glue=None):
    # 模拟物料的规格串。核对时用同一格式，避免 25 和 25.0 对不上。灌胶取自核算模板面别。
    face = glue or "单面"
    return f"P{_pitch(row['点间距_mm'])} {int(float(row['模组宽_mm']))}x{int(float(row['模组高_mm']))} {row.get('配比')} {face}"


def box_spec(row):
    return f"{row.get('材料') or ''} {int(float(row['宽_mm']))}x{int(float(row['高_mm']))}"
