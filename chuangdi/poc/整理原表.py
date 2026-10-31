# 把客户原表整理成：规格单价表（材料成本）+ 核算口径表（大模组核算）。
# 规格必转；单价：原表有数字就转入，无数留空（不填 0）。供应商名称不转入。原件留在 docs，不覆盖。

import re
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side

ROOT = Path(__file__).resolve().parent
CHUANGDI = ROOT.parent
SRC_BOOK = CHUANGDI / "docs" / "材料成本.xlsx"
SRC_COST = CHUANGDI / "docs" / "大模组核算2026.06.16.xls"
TPL = CHUANGDI / "模版" / "报价数据整理模版.xlsx"
OUT = CHUANGDI / "数据"
# 整理结果与模版约定一致：同列同色同字号。蓝=手填，绿=公式。备注列一律留空。
_FONT_IN = Font(name="微软雅黑", size=10, color="0000FF")
_FONT_FX = Font(name="微软雅黑", size=10, color="008000")
_NO_FILL = PatternFill(fill_type=None)
_EDGE = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def paint(cell, value, kind="in"):
    # kind: in=手填，fx=公式，note=备注（仍蓝，与手填同列观感）
    cell.value = value
    cell.font = _FONT_FX if kind == "fx" else _FONT_IN
    cell.fill = _NO_FILL
    cell.border = _EDGE


def count_of(value):
    # 灯珠颗数在原表里有的是数，有的是乘法公式。只计算数字运算，不算单价。
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str) or not value.startswith("="):
        return None
    expr = value[1:].replace(" ", "")
    if not re.fullmatch(r"[\d.+\-*/()]+", expr):
        return None
    try:
        return float(eval(expr, {"__builtins__": {}}, {}))
    except Exception:
        return None


def pitch_of(value):
    if value is None:
        return None
    text = str(value).strip()
    if text[:1] in ("P", "p"):
        text = text[1:]
    text = text.replace("贴片", "").replace("直插", "")
    try:
        return float(text)
    except ValueError:
        return None


def size_of(value):
    if value is None:
        return None, None
    found = re.search(r"([\d.]+)\s*[*×xX]\s*([\d.]+)", str(value))
    if not found:
        return None, None
    return float(found.group(1)), float(found.group(2))


def bright_of(value):
    if value is None:
        return None
    found = re.search(r"(\d+)", str(value))
    return int(found.group(1)) if found else None


def color_of(value):
    text = str(value or "")
    if "黄" in text and "白" in text:
        return "黄白"
    if "黄" in text:
        return "黄色"
    if "白" in text:
        return "白色"
    return text[:12] or None


def batch_of(formula):
    text = str(formula or "")
    if "A3" in text:
        return "批次-510"
    if "A7" in text:
        return "批次-310"
    return None


def blocks(ws):
    # 同一张表左右两块是两个灯珠批次。右侧没有名称就只有一块。
    left = batch_of(ws.cell(3, 12).value)
    right = batch_of(ws.cell(3, 31).value)
    out = [(2, left or "批次-未标明")]
    if ws.cell(4, 21).value:
        out.append((21, right or "批次-未标明"))
    return out


def section_of(ws, start, row):
    # 同列下方可能再挂一段标题，如「普通模组价格5124-310-F5」；向上找最近标题。
    # 标题在名称列左侧一格（序号列）或本列。
    for r in range(row, 0, -1):
        for cell in (ws.cell(r, start - 1).value, ws.cell(r, start).value):
            text = str(cell or "")
            if "模组价格" in text and "-" in text:
                return text.strip()
    return None


def chip_of(value, section):
    # 驱动芯片统一成文本，避免 5124 被写成数字后筛选看不到。
    if value not in (None, ""):
        return str(value).strip()
    if not section:
        return None
    # 标题形如 普通模组价格5124-310-F5 / 250*500模组价格5124-310-F5
    found = re.search(r"模组价格([A-Za-z0-9]+)-", section)
    return found.group(1) if found else None


def money(value):
    # 只收数字单价；公式或空不转成 0。
    if isinstance(value, (int, float)):
        return float(value)
    return None


def glue_of(double, single=None):
    # 双面公式常见 =单面单元格+常数；或两列都是数字时用差。
    if isinstance(double, (int, float)) and isinstance(single, (int, float)):
        return float(double) - float(single)
    if isinstance(double, str) and double.startswith("="):
        found = re.search(r"\+(\d+(?:\.\d+)?)\s*$", double.replace(" ", ""))
        if found:
            return float(found.group(1))
    return None


def col_map(ws, start):
    # 按第 2 行表头找加工费/单面/双面列。250-500 表在加工费前多一列套件单价。
    out = {}
    for c in range(start, start + 20):
        head = str(ws.cell(2, c).value or "")
        if "加工费" in head:
            out["fee"] = c
        elif "单面" in head:
            out["single"] = c
        elif "双面" in head:
            out["double"] = c
    return out


def lamps(wb):
    # 灯价格表：有数字的 R/G/B 才转入。供应商名不读。
    ws = wb["灯价格"]
    rows = []
    for batch, r in (("批次-510", 3), ("批次-310", 7)):
        rows.append({
            "批次": batch,
            "R": money(ws.cell(r, 1).value),
            "G": money(ws.cell(r, 2).value),
            "B": money(ws.cell(r, 3).value),
        })
    return rows


def modules(wb):
    rows = []
    sheets = ["普通模组价格5034B", "普通模组价格16289", "250-500模组价格5034", "250-500模组价格16289"]
    for name in sheets:
        ws = wb[name]
        for start, batch in blocks(ws):
            cols = col_map(ws, start)
            for r in range(4, (ws.max_row or 4) + 1):
                title = ws.cell(r, start).value
                if pitch_of(title) is None:
                    continue
                # 原表同行两列含税单价（单面/双面）保留在同一规格行，报价时按核算模板「灌胶」选列。
                # 同表下方还可再挂 5124 等驱动分段（如 5034B 表右侧第 32 行起）。
                wide, high = size_of(ws.cell(r, start + 4).value)
                section = section_of(ws, start, r)
                fee_c = cols.get("fee")
                single_c = cols.get("single")
                double_c = cols.get("double")
                single = ws.cell(r, single_c).value if single_c else None
                double = ws.cell(r, double_c).value if double_c else None
                rows.append({
                    "名称": str(title).strip(),
                    "点间距_mm": pitch_of(title),
                    "配比": ws.cell(r, start + 1).value,
                    "灯珠芯片": ws.cell(r, start + 2).value,
                    "驱动芯片": chip_of(ws.cell(r, start + 5).value, section),
                    "亮度": bright_of(ws.cell(r, start + 3).value),
                    "宽": wide,
                    "高": high,
                    "色品": color_of(ws.cell(r, start + 6).value),
                    "批次": batch,
                    "工艺": "贴片",
                    "R": count_of(ws.cell(r, start + 7).value),
                    "G": count_of(ws.cell(r, start + 8).value),
                    "B": count_of(ws.cell(r, start + 9).value),
                    "加工费": money(ws.cell(r, fee_c).value) if fee_c else None,
                    "灌胶加价": glue_of(double, single),
                    "来源": name,
                    "分区": section,
                    "原表行": r,
                })
    return rows


def kits(wb):
    ws = wb["套件"]
    rows = []
    for r in range(3, 20):
        title = ws.cell(r, 1).value
        if title is None:
            continue
        text = str(title)
        craft = "贴片" if "贴片" in text else "直插" if "直插" in text else None
        pitch = pitch_of(text)
        if pitch is None:
            continue
        rows.append({
            "点间距_mm": pitch,
            "工艺": craft,
            "底壳": money(ws.cell(r, 2).value),
            "面罩数量": money(ws.cell(r, 3).value),
            "面罩单价": money(ws.cell(r, 4).value),
            "胶圈": money(ws.cell(r, 6).value),
            "来源": f"套件第{r}行",
        })
    return rows


def boxes(wb):
    ws = wb["箱体单价"]
    rows = []
    for r in range(3, 14):
        title = ws.cell(r, 2).value
        if not title:
            continue
        note = str(ws.cell(r, 5).value or "")
        seal = "后开放" if "后开放" in note else "全密封" if "全密封" in note else None
        edge = "含" if "含包边" in note else "不含" if "不含包边" in note else None
        wide, high = size_of(str(ws.cell(r, 3).value or "").replace("*", "*"))
        # 原表尺寸是米。模版这一列是毫米。
        if wide is not None and wide < 20:
            wide, high = wide * 1000, high * 1000
        place = "洞内" if "洞内" in str(title) else "门架" if "门架" in str(title) else "悬臂" if "悬臂" in str(title) else "车道" if "车道" in str(title) else None
        material = "铁箱" if "铁" in str(title) else "压铸铝" if "铝" in str(title) else None
        rows.append({
            "名称": f"{title}-{seal or '密封未写'}-{edge or '包边未写'}",
            "材料": material,
            "位置": place,
            "宽": wide,
            "高": high,
            "密封": "后开放" if seal == "后开放" else "全密封" if seal == "全密封" else None,
            "包边": edge,
            "含税单价": money(ws.cell(r, 4).value),
        })
    return rows


def shift(formula, src, dest):
    return re.sub(rf"([A-Z]){src}(?!\d)", rf"\g<1>{dest}", formula)


def write_price(mods, kit_rows, box_rows, lamp_rows):
    dest = OUT / "规格单价表.xlsx"
    shutil.copy(TPL, dest)
    wb = load_workbook(dest)
    if "供应商" in wb.sheetnames:
        del wb["供应商"]
    wb["说明"]["A1"] = (
        "本文件按模版整理自材料成本.xlsx。原表有数字的单价已转入，无数留空。"
        "不含供应商表。表头在第 4 行，数据从第 5 行起。"
    )
    # 列布局见模版：无「是否当前」「生效日期」。写入时统一字体，避免同列蓝/黑、10/11 混用。
    lamp = wb["灯珠单价"]
    for i, row in enumerate(lamp_rows, start=5):
        paint(lamp[f"A{i}"], row["批次"])
        paint(lamp[f"B{i}"], row.get("R"))
        paint(lamp[f"C{i}"], row.get("G"))
        paint(lamp[f"D{i}"], row.get("B"))
        paint(lamp[f"E{i}"], None)  # 备注清空
    kit = wb["套件单价"]
    src_g = kit["G5"].value
    for i, row in enumerate(kit_rows, start=5):
        paint(kit[f"A{i}"], row["点间距_mm"])
        paint(kit[f"B{i}"], row["工艺"])
        paint(kit[f"C{i}"], row.get("底壳"))
        paint(kit[f"D{i}"], row.get("面罩数量"))
        paint(kit[f"E{i}"], row.get("面罩单价"))
        paint(kit[f"F{i}"], row.get("胶圈"))
        paint(kit[f"G{i}"], shift(src_g, 5, i), "fx")
        paint(kit[f"H{i}"], None)  # 备注清空
    mod = wb["模组价格"]
    # A名称 … J批次 K工艺 L/M/N数量 O/P/Q灯珠单价 R合计 S加工费 T套件 U灌胶加价 V单面 W双面 X备注
    formulas = {col: mod[f"{col}5"].value for col in ("O", "P", "Q", "R", "T", "V", "W")}
    for i, row in enumerate(mods, start=5):
        paint(mod[f"A{i}"], row["名称"])
        paint(mod[f"B{i}"], row["点间距_mm"])
        paint(mod[f"C{i}"], row["配比"])
        paint(mod[f"D{i}"], row["灯珠芯片"])
        paint(mod[f"E{i}"], row["驱动芯片"])
        paint(mod[f"F{i}"], row["亮度"])
        paint(mod[f"G{i}"], row["宽"])
        paint(mod[f"H{i}"], row["高"])
        paint(mod[f"I{i}"], row["色品"])
        paint(mod[f"J{i}"], row["批次"])
        paint(mod[f"K{i}"], row["工艺"])
        paint(mod[f"L{i}"], row["R"])
        paint(mod[f"M{i}"], row["G"])
        paint(mod[f"N{i}"], row["B"])
        paint(mod[f"S{i}"], row.get("加工费"))
        paint(mod[f"U{i}"], row.get("灌胶加价"))
        for col, formula in formulas.items():
            paint(mod[f"{col}{i}"], shift(formula, 5, i), "fx")
        paint(mod[f"X{i}"], None)  # 备注清空
    box = wb["箱体单价"]
    for i, row in enumerate(box_rows, start=5):
        paint(box[f"A{i}"], row["名称"])
        paint(box[f"B{i}"], row["材料"])
        paint(box[f"C{i}"], row["位置"])
        paint(box[f"D{i}"], row["宽"])
        paint(box[f"E{i}"], row["高"])
        paint(box[f"F{i}"], row["密封"])
        paint(box[f"G{i}"], row["包边"])
        paint(box[f"H{i}"], row.get("含税单价"))
        paint(box[f"I{i}"], None)  # 备注清空
    calc = wb["核算模板"]
    for col in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"):
        paint(calc[f"{col}5"], None)
        paint(calc[f"{col}6"], None)
    paint(calc["T5"], None)  # 备注清空
    paint(calc["T6"], None)
    # 核算模板绿色勾稽公式列也统一字号
    for col in ("R", "S"):
        for i in (5, 6):
            cell = calc[f"{col}{i}"]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                paint(cell, cell.value, "fx")
    wb.save(dest)
    return dest


def _xlsx_src():
    # 原核算是 xls；转成 xlsx 后保留公式再整理。
    folder = Path("/tmp/lo-hesuan")
    folder.mkdir(exist_ok=True)
    xlsx = folder / "大模组核算2026.06.16.xlsx"
    if not xlsx.exists():
        subprocess.check_call(
            ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(folder), str(SRC_COST)]
        )
    return xlsx


def _copy_sheet(src, dst, remap):
    # 按已用区域抄单元格；公式里的工作表名按 remap 替换。
    for row in src.iter_rows(min_row=1, max_row=src.max_row or 1, max_col=min(src.max_column or 1, 20)):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and val.startswith("="):
                for old, new in remap.items():
                    val = val.replace(old, new)
            dst.cell(cell.row, cell.column, val)


def _rewrite_fx(val, remap):
    if not isinstance(val, str) or not val.startswith("="):
        return val
    for old, new in remap.items():
        val = val.replace(old, new)
    return val


def cost_book():
    # 核算口径表：辅材 BOM + 每平米收口骨架 + 箱体/电源/支架/抱箍价目。
    # 模组单价单源在规格单价表；箱体进平方单价用本册「箱体」元/㎡价目（与规格表物理箱 SKU 分工）。
    src = load_workbook(_xlsx_src(), data_only=False)
    out = Workbook()
    tip = out.active
    tip.title = "说明"
    tip["A1"] = (
        "本文件整理自大模组核算原表：辅材 BOM、每平米收口结构、箱体/电源/支架/抱箍/控制箱价目。"
        "「每平米收口」样例行是历史底稿，项目报价由 Agent 按招标+规格+本表规则重算本项目收口，不是样例×面积。"
    )
    tip["A2"] = (
        "三件套：规格单价表（模组/灯珠/套件及物理箱 SKU）+ 本表（BOM、收口结构、元/㎡箱体与电源/支架/抱箍/控制箱价目）"
        "+ 计价 skill。详见方案/场景演示建议口径.md、原表公式与经验配量确认.md。"
    )
    tip["A3"] = "压铸铝辅材 / 铁箱辅材：保留原逐项与汇总公式。关键汇总格见各表「总计1平方辅料」「包材总计」「折1平方」行。"
    tip["A4"] = (
        "每平米收口：保留结构与对辅材/箱体/电源的公式引用；模组单价列正式稿留空（取规格单价表）。"
        "表内「1㎡模组数量」「常规电源数量」为样例/待成本确认的经验位，不作未确认的项目默认取数。"
    )
    tip["A5"] = (
        "价目页「箱体」「电源」「支架」「安装抱箍和螺钉」「控制箱」自原表迁入；正式稿单价空/0 由成本维护。"
        "规格单价表「箱体单价」仅作物料核对，平方单价中的箱体元/㎡以本册「箱体」为准。"
        "报价单「控制系统」取本册「控制箱」汇总（每套），形态加价见报价单 skill。"
    )
    tip["A6"] = "不转入屏体清单样例页。招标用抄录面积。供应商名不转入。"
    tip["A7"] = "BOM 明细与常数加价（如辅材 +20）由成本岗维护，改后同步 skill。"

    remap = {
        "1压铸铝平方辅材": "压铸铝辅材",
        "1铁箱平方辅材": "铁箱辅材",
    }
    al = out.create_sheet("压铸铝辅材")
    _copy_sheet(src["1压铸铝平方辅材"], al, remap)
    fe = out.create_sheet("铁箱辅材")
    _copy_sheet(src["1铁箱平方辅材"], fe, remap)

    for name in ("箱体", "电源", "支架", "安装抱箍和螺钉"):
        dst = out.create_sheet(name)
        _copy_sheet(src[name], dst, {})

    # 原表「压铸铝控制箱」→ 工作表「控制箱」，供报价单控制系统取汇总
    ctrl = out.create_sheet("控制箱")
    _copy_sheet(src["压铸铝控制箱"], ctrl, {})

    raw = src["1平方大模组屏体核算"]
    sqm = out.create_sheet("每平米收口")
    n_data = 0
    for r in range(1, (raw.max_row or 1) + 1):
        for c in range(1, 18):
            val = raw.cell(r, c).value
            val = _rewrite_fx(val, remap)
            # 模组单价列：原表手填数字不转入，避免与规格单价表双源。
            if c == 6 and isinstance(val, (int, float)):
                val = None
            sqm.cell(r, c, val)
        a = raw.cell(r, 1).value
        if a and re.match(r"^\d", str(a)):
            n_data += 1
            note = sqm.cell(r, 17).value
            mark = "模组单价见规格单价表；箱体/电源见本册价目"
            sqm.cell(r, 17, f"{note}；{mark}" if note else mark)

    path = OUT / "核算口径表.xlsx"
    out.save(path)
    return path, n_data


def read_log(n_mod, n_kit, n_box, n_cost, _n_src=None):
    text = f"""# 整理时怎么读的原表

原件不进入云端任务。这里记录整理时读了哪些表、表头在哪一行、哪些列被留下。

## 材料成本.xlsx → 规格单价表.xlsx

| 工作表 | 表头行 | 原表列（相对每块起点） | 写入「模组价格」 |
| --- | --- | --- | --- |
| 普通模组价格5034B、普通模组价格16289 | 第 2 行 | 名称、配比、芯片、亮度、尺寸、聚积芯片、色品、R/G/B 数量。左块起点列 2（批次-510），右块起点列 21（批次-310，有名称才读）。同列下方可再挂分段标题（如 5124-310-F5） | 一行原规格 → 一行；驱动芯片统一文本；单面/双面含税单价同属该行两列 |
| 250-500模组价格5034、250-500模组价格16289 | 第 2 行 | 同上；仅左块，批次-310；下方可挂 5124 分段 | 同上 |
| 灯价格 | 第 2、6 行是单价表头 | 有数字的 R/G/B 转入；供应商名不读 | 灯珠单价表批次-510、批次-310 |
| 套件 | 第 1–2 行 | 名称中的点间距与贴片/直插；有数字的底壳/面罩/胶圈转入 | 套件单价表 |
| 箱体单价 | 第 2 行 | 名称、尺寸、备注中的密封与包边；有数字的单价转入 | 箱体单价表；同名按密封/包边拆开 |
| 折叠屏模组、强力 | — | 无规格行 | 不转入 |

原表模组规格行整理后写入 {n_mod} 行（与原规格一一对应，单面/双面不拆行）。套件 {n_kit} 行，箱体 {n_box} 行。表头在第 4 行。无「供应商」工作表。

列对照见 `模版/填写说明.md`。

## 大模组核算2026.06.16.xls → 核算口径表.xlsx

先转 xlsx，**保留公式**再整理。多工作表：

| 工作表 | 来源 | 留下的内容 |
| --- | --- | --- |
| 说明 | — | 三件套分工；本项目收口由 Agent 重算，非样例×面积 |
| 压铸铝辅材 | 1压铸铝平方辅材 | 逐项与汇总公式；`F56`/`F60`/`E80` 等 |
| 铁箱辅材 | 1铁箱平方辅材 | 同上；`E58`/`E63` 等 |
| 箱体 / 电源 / 支架 / 安装抱箍和螺钉 | 同名原表 | 价目结构；正式稿单价空/0 由成本维护 |
| 控制箱 | 压铸铝控制箱 | BOM；合计列汇总供报价单「控制系统」每套取值 |
| 每平米收口 | 1平方大模组屏体核算 | 结构与公式引用；模组单价列留空。数据行约 {n_cost} 行 |

不转入：屏体清单样例层。原公式链见 `方案/核算原表计算规则分析.md`。
"""
    (OUT / "读表记录.md").write_text(text, encoding="utf-8")


def main():
    OUT.mkdir(exist_ok=True)
    src = load_workbook(SRC_BOOK, data_only=False)
    mods = modules(src)
    kit_rows = kits(src)
    box_rows = boxes(src)
    lamp_rows = lamps(src)
    write_price(mods, kit_rows, box_rows, lamp_rows)
    _, n_cost = cost_book()
    read_log(len(mods), len(kit_rows), len(box_rows), n_cost, len(mods))
    blob = (OUT / "规格单价表.xlsx").read_bytes() + (OUT / "核算口径表.xlsx").read_bytes()
    for word in ("弘正", "鑫恩拓"):
        if word.encode("utf-8") in blob:
            raise SystemExit(f"整理结果含有不应出现的名称：{word}")
    print(f"模组 {len(mods)} 套件 {len(kit_rows)} 箱体 {len(box_rows)} 核算口径 {n_cost}")
    print(OUT / "规格单价表.xlsx")
    print(OUT / "核算口径表.xlsx")


if __name__ == "__main__":
    main()
