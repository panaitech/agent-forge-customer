# 计价岗读表。只读本次任务提交目录里的工作簿，并写下表头和列。
# 不按「是否当前」过滤（模版已取消该列）。不重算金额。

from openpyxl import load_workbook
from pathlib import Path


def _sheet(path, name, head_row):
    wb = load_workbook(path, data_only=False)
    if "供应商" in wb.sheetnames:
        raise SystemExit(f"{path.name} 含有供应商表，本次任务退回")
    ws = wb[name]
    head = []
    for cell in ws[head_row]:
        if cell.value:
            head.append(str(cell.value))
    count = 0
    empty_price = 0
    price_cols = {"R单价", "G单价", "B单价", "加工费", "灌胶加价", "含税单价_单面", "含税单价_双面", "底壳单价", "金额", "含税单价"}
    for raw in ws.iter_rows(min_row=head_row + 1, values_only=True):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        item = {}
        for i, key in enumerate(head):
            item[key] = raw[i] if i < len(raw) else None
        if all(item.get(k) in (None, "") for k in item):
            continue
        count += 1
        if any(k in item and item.get(k) in (None, "") for k in price_cols if k in head):
            empty_price += 1
    return head, count, empty_price


def report(folder: Path):
    lines = ["# 读表", "", "只读本次任务提交目录。没有去任务以外的路径。", ""]
    for path in sorted(folder.glob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=False)
        names = wb.sheetnames
        wb.close()
        lines.append(f"## {path.name}")
        lines.append("")
        lines.append("工作表：" + "、".join(names))
        if path.name == "核算口径表.xlsx":
            for sheet in ("压铸铝辅材", "铁箱辅材", "每平米收口"):
                if sheet not in names:
                    continue
                # BOM 表头第 1 行；每平米收口表头在各点间距块内（首块第 2 行）
                head_row = 2 if sheet == "每平米收口" else 1
                head, count, empty = _sheet(path, sheet, head_row)
                lines.append(f"### {sheet}")
                lines.append(f"表头在第 {head_row} 行。列：" + "、".join(head[:12]) + ("…" if len(head) > 12 else ""))
                lines.append(f"非空数据约 {count} 行。")
                if sheet == "压铸铝辅材":
                    lines.append("汇总格关注 F56（1㎡辅料）、F60（包材）、E80（10㎡折1㎡）。")
                if sheet == "铁箱辅材":
                    lines.append("汇总格关注 E58（1㎡辅料）、E63（包材）。")
                if sheet == "每平米收口":
                    lines.append("数量口径；模组/箱体/电源单价列应留空。包材/辅材公式引用辅材表。")
        elif "模组价格" in names:
            head, count, empty = _sheet(path, "模组价格", 4)
            lines.append("模组价格表头在第 4 行。列：" + "、".join(head))
            lines.append(f"模组 {count} 行。单价列有空值的行 {empty} 行。")
        lines.append("")
    return "\n".join(lines)
