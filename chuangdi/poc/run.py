# 本地工具：价表 selfcheck；可选 --pdf 试跑（产物目录由 --out 指定）。
# 不是 Agent Forge 协作入口。无 MCP。

import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
CHUANGDI = ROOT.parent

import bill
import book
import erp
import readbook
import 整理原表


def submit(task, pdf, instruction):
    folder = task / "提交"
    if folder.exists():
        shutil.rmtree(folder)
    folder.mkdir(parents=True)
    # 正式整理稿（单价可空）+ 演示稿（模拟价，非客户成本）
    for name in ("规格单价表.xlsx", "核算口径表.xlsx", "规格单价表-演示.xlsx", "核算口径表-演示.xlsx"):
        src = CHUANGDI / "数据" / name
        if not src.exists():
            raise SystemExit(f"缺少 {src}，请先运行整理原表.py 与 seed_demo.py")
        shutil.copy(src, folder / name)
    shutil.copy(instruction, folder / "解析指令.json")
    shutil.copy(pdf, folder / pdf.name)
    return folder


def parse_task(folder):
    pdf = next(folder.glob("*.pdf"))
    cfg = bill.load(folder / "解析指令.json")
    found = bill.pages(pdf)
    rows = bill.parse(
        "\n".join(found[n] for n in cfg["pdf_printed_pages"]["list"]),
        "\n".join(found[n] for n in cfg["pdf_printed_pages"]["spec"]),
        cfg,
    )
    lines = ["# 解析岗", "", "依据 scripts/skills/chuangdi-quote-parse。不读价格表。", ""]
    for row in rows:
        if row.get("skip"):
            lines.append(f"- 排除 {row['name']}")
            continue
        if not row.get("found"):
            lines.append(f"- 未找到 {row['name']}")
            continue
        bits = [row["name"], f"数量 {row.get('qty')}{row.get('unit') or ''}"]
        if row.get("w"):
            bits.append(f"面积 {row['w']}×{row['h']} 米，每套 {row.get('panels')} 块")
        if row.get("pitch"):
            bits.append(f"点间距 {row['pitch']} mm，{row.get('pitch_source')}")
        else:
            bits.append("无毫米点间距")
        lines.append("- " + "；".join(bits))
    return rows, "\n".join(lines) + "\n"


def review(rows):
    lines = ["# 复核岗", "", "依据 scripts/skills/chuangdi-quote-review。不改抄录，不取价。", ""]
    for row in rows:
        if row.get("skip") or not row.get("found"):
            continue
        if row.get("pitch") and row.get("pitch_source"):
            lines.append(f"- {row['name']}：出处页能对上点间距 {row['pitch']} mm，抄录通过。")
        elif row.get("kind") == "screen" and not row.get("pitch"):
            lines.append(f"- {row['name']}：抄录通过。评分页没有毫米点间距，后面的岗不取模组价。")
        else:
            lines.append(f"- {row['name']}：抄录通过。{'；'.join(row.get('reason') or [])}")
    return "\n".join(lines) + "\n"


def quote_hand(rows):
    # 只整理抄录与文件指针；金额由 Agent 按 skill 计算，程序不重算。
    lines = [
        "# 计价岗",
        "",
        "依据 scripts/skills/chuangdi-quote-price。",
        "",
        "输入三件套：规格单价表（大件单价）+ 核算口径表（辅材 BOM 与每平米收口）+ 本 skill。",
        "正式稿单价空则金额留空。需要非零对照时用「规格单价表-演示.xlsx」「核算口径表-演示.xlsx」（模拟价，非客户成本）。",
        "本程序不另写一套计价算法。",
        "",
        "## 抄录行（供 skill）",
        "",
    ]
    for row in rows:
        if row.get("skip") or not row.get("found"):
            continue
        if row.get("kind") != "screen":
            lines.append(f"- {row['name']}：非屏行。{'；'.join(row.get('reason') or [])}")
            continue
        bits = [row["name"]]
        if row.get("w") is not None:
            bits.append(f"面积 {row['w']}×{row['h']} 米")
        bits.append(f"每套 {row.get('panels') or 1} 块，套数 {row.get('qty')}")
        if row.get("pitch"):
            bits.append(f"点间距 {row['pitch']} mm")
        else:
            bits.append("无毫米点间距，不取模组价，金额留空")
        if row.get("install"):
            bits.append(f"安装形态 {row['install']}")
        if row.get("reason"):
            bits.append("；".join(row["reason"]))
        lines.append("- " + "；".join(bits))
    lines.extend([
        "",
        "## 演示稿",
        "",
        "规格单价表-演示.xlsx、核算口径表-演示.xlsx 仅供按 skill 对照练习，不能写成正式报价。本交接不由程序重算金额。",
        "",
    ])
    return "\n".join(lines)


def land(pdf, instruction, task):
    if not (CHUANGDI / "数据" / "规格单价表.xlsx").exists():
        整理原表.main()
    if not (CHUANGDI / "数据" / "规格单价表-演示.xlsx").exists():
        import seed_demo
        seed_demo.main()
    folder = submit(task, pdf, instruction)
    hand = task / "交接"
    hand.mkdir(exist_ok=True)
    rows, text = parse_task(folder)
    (hand / "01-解析.md").write_text(text, encoding="utf-8")
    (hand / "02-复核.md").write_text(review(rows), encoding="utf-8")
    (hand / "03-读表.md").write_text(readbook.report(folder), encoding="utf-8")
    data = erp.seed(book.load(folder / "规格单价表-演示.xlsx"))
    (hand / "04-计价.md").write_text(quote_hand(rows), encoding="utf-8")
    sample = erp.query(data, "BD_MATERIAL", "FNumber,FName,FSpecification,FForbidStatus", "FName='模组'")
    # 故意用不存在的规格演示「不一致」：状态外附带同名候选，供销售对照
    mismatch = erp.check_part(data, "模组", "P25 999x999 演示核对规格")
    (hand / "05-物料核对.md").write_text(
        "# 物料核对\n\n对照附件规格单价表（本地 erp.check_part 示意）。Forge 协作由岗位只读价表核对，无 MCP。\n\n"
        "返回仅编码/名称/规格/禁用状态，不含单价与采购来源。不一致/禁用时附带候选 hits。\n\n"
        "## 列表查询\n\n```\n"
        + json.dumps(sample[:8], ensure_ascii=False, indent=2)
        + ("\n  …\n" if len(sample) > 8 else "\n")
        + "```\n\n## 核对示例（规格对不上 → 不一致 + hits）\n\n```\n"
        + json.dumps(mismatch, ensure_ascii=False, indent=2)
        + "\n```\n",
        encoding="utf-8",
    )
    sheet_lines = [
        "# 报价单",
        "",
        "| 项目名称 | （取自招标附件标题） |",
        "| 招标编号 | （有则写） |",
        "| 报价日期 | （试跑占位） |",
        "| 备注 | 非正式签署件；空价待销售确认；演示价未计入正式栏 |",
        "",
        "## 报价明细",
        "",
        "| 序号 | 招标清单名称 | 规格摘要 | 数量 | 单位 | 单价（含税） | 金额（含税） | 出处 | 备注 |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    n = 0
    for row in rows:
        if row.get("skip") or not row.get("found"):
            continue
        if row.get("kind") != "screen":
            continue
        n += 1
        pitch = f"P{row['pitch']}" if row.get("pitch") else "待确认"
        sheet_lines.append(
            f"| {n} | {row['name']} | {pitch} | {row.get('qty') or ''} | {row.get('unit') or ''} | 待确认 | 待确认 | {row.get('pitch_source') or '见解析'} | 正式单价空则待确认 |"
        )
    if n == 0:
        sheet_lines.append("| 1 | （无复核通过的屏行） |  |  |  | 待确认 | 待确认 |  |  |")
    sheet_lines.extend(
        [
            "",
            "## 合计",
            "",
            "| 已报价小计 | 待确认 |",
            f"| 待确认行数 | {n} |",
            "",
            "本单不改招标数量。演示价未计入正式栏。不定责、不盖章、不建单。",
            "",
        ]
    )
    md = "\n".join(sheet_lines)
    (hand / "报价单.md").write_text(md, encoding="utf-8")
    return hand


def main():
    import argparse

    parser = argparse.ArgumentParser(description="价表 selfcheck，或本地试跑（须指定 --out）")
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--instruction", type=Path, help="本地试跑用的解析说明 JSON；Forge 在任务说明里写")
    parser.add_argument("--out", type=Path, help="本地试跑工作区（提交/交接），须显式指定")
    parser.add_argument("--selfcheck", action="store_true")
    args = parser.parse_args()
    if args.selfcheck:
        selfcheck()
        return
    if not args.pdf or not args.instruction or not args.out:
        parser.error("试跑须同时传入 --pdf、--instruction、--out；或使用 --selfcheck")
    hand = land(args.pdf, args.instruction, args.out)
    print(hand)


def selfcheck():
    """只校验本机整理出的价表/演示稿；完整报价链在 Agent Forge 跑。"""
    整理原表.main()
    import seed_demo
    seed_demo.main()
    from openpyxl import load_workbook

    wb = load_workbook(CHUANGDI / "数据" / "规格单价表.xlsx", read_only=True)
    if "供应商" in wb.sheetnames:
        raise SystemExit("规格单价表不应含供应商")
    wb.close()
    cost = load_workbook(CHUANGDI / "数据" / "核算口径表.xlsx", data_only=False)
    need = {"说明", "压铸铝辅材", "铁箱辅材", "每平米收口"}
    if not need.issubset(set(cost.sheetnames)):
        raise SystemExit(f"核算口径表缺少工作表：{need - set(cost.sheetnames)}")
    al = cost["压铸铝辅材"]
    if not (isinstance(al["F56"].value, str) and al["F56"].value.startswith("=")):
        raise SystemExit("压铸铝辅材 F56 应保留汇总公式")
    sqm = cost["每平米收口"]
    o3 = sqm["O3"].value
    if not (isinstance(o3, str) and "压铸铝辅材" in o3):
        raise SystemExit("每平米收口辅材公式应引用压铸铝辅材")
    if sqm["F3"].value not in (None, ""):
        raise SystemExit("正式核算口径表模组单价应留空（单价单源在规格单价表）")
    cost.close()
    demo_u = load_workbook(CHUANGDI / "数据" / "规格单价表-演示.xlsx", data_only=False)
    if demo_u["灯珠单价"]["B5"].value in (None, ""):
        raise SystemExit("规格单价表-演示应填入灯珠单价")
    if demo_u["箱体单价"]["H5"].value in (None, ""):
        raise SystemExit("规格单价表-演示应填入箱体单价")
    demo_u.close()
    demo_c = load_workbook(CHUANGDI / "数据" / "核算口径表-演示.xlsx", data_only=False)
    if demo_c["每平米收口"]["F3"].value in (None, ""):
        raise SystemExit("核算口径表-演示应收口模组单价有示意数")
    if demo_c["压铸铝辅材"]["D9"].value in (None, ""):
        raise SystemExit("核算口径表-演示辅材应有模拟单价")
    demo_c.close()
    blob = (CHUANGDI / "数据" / "规格单价表.xlsx").read_bytes() + (CHUANGDI / "数据" / "核算口径表.xlsx").read_bytes()
    for word in ("弘正", "鑫恩拓"):
        if word.encode() in blob:
            raise SystemExit(word)
    print("selfcheck ok")


if __name__ == "__main__":
    main()
